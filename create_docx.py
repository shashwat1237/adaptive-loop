import os
import shutil
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

DOCS_DIR = r"c:\Users\admin\Downloads\broyo\docs"
SCREENSHOTS_DIR = os.path.join(DOCS_DIR, "screenshots")
OUTPUT_DOCX_DOCS = os.path.join(DOCS_DIR, "Adaptive_Loop_Demo_Walkthrough.docx")
ARTIFACT_DIR = r"C:\Users\admin\.gemini\antigravity\brain\f31d380e-4922-4144-93ec-efbe752ebaf3"
OUTPUT_DOCX_ARTIFACT = os.path.join(ARTIFACT_DIR, "Adaptive_Loop_Demo_Walkthrough.docx")

def set_cell_background(cell, hex_color):
    """Sets the background color of a table cell."""
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), hex_color)
    tcPr.append(shd)

def set_cell_margins(cell, top=140, bottom=140, left=180, right=180):
    """Sets internal padding for a table cell in twips."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for margin_name, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{margin_name}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def add_callout(doc, text, title=None, hex_bg="F1F5F9", hex_border="0284C7"):
    """Adds a stylish callout box with a colored background."""
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    tbl.columns[0].width = Inches(6.5)
    cell = tbl.cell(0, 0)
    set_cell_background(cell, hex_bg)
    set_cell_margins(cell, top=160, bottom=160, left=220, right=220)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.15
    if title:
        run_title = p.add_run(f"📌 {title}\n")
        run_title.bold = True
        run_title.font.name = "Arial"
        run_title.font.size = Pt(10.5)
        run_title.font.color.rgb = RGBColor(2, 132, 199)
    run_text = p.add_run(text)
    run_text.font.name = "Arial"
    run_text.font.size = Pt(10)
    run_text.font.color.rgb = RGBColor(30, 41, 59)
    
    # Empty space after callout
    p_spacer = doc.add_paragraph()
    p_spacer.paragraph_format.space_before = Pt(0)
    p_spacer.paragraph_format.space_after = Pt(4)

def add_screenshot_figure(doc, filename, caption):
    """Embeds a screenshot into the Word document with proper width and caption."""
    img_path = os.path.join(SCREENSHOTS_DIR, filename)
    if not os.path.exists(img_path):
        print(f"Warning: Image not found: {img_path}")
        return
    
    p_img = doc.add_paragraph()
    p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_img.paragraph_format.space_before = Pt(8)
    p_img.paragraph_format.space_after = Pt(4)
    run_img = p_img.add_run()
    run_img.add_picture(img_path, width=Inches(6.4))
    
    p_caption = doc.add_paragraph()
    p_caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_caption.paragraph_format.space_before = Pt(2)
    p_caption.paragraph_format.space_after = Pt(14)
    run_cap = p_caption.add_run(caption)
    run_cap.italic = True
    run_cap.font.name = "Arial"
    run_cap.font.size = Pt(9.5)
    run_cap.font.color.rgb = RGBColor(100, 116, 139)

def build_docx():
    print("Building comprehensive Word document with embedded screenshots...")
    doc = docx.Document()
    
    # 1. Page Setup (0.75 in margins for modern presentation)
    for section in doc.sections:
        section.top_margin = Inches(0.75)
        section.bottom_margin = Inches(0.75)
        section.left_margin = Inches(0.75)
        section.right_margin = Inches(0.75)
        
        # Header / Footer
        footer = section.footer
        p_ft = footer.paragraphs[0]
        p_ft.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        r_ft = p_ft.add_run("Adaptive Loop • Clinical Decision Support • System Walkthrough")
        r_ft.font.name = "Arial"
        r_ft.font.size = Pt(8.5)
        r_ft.font.color.rgb = RGBColor(148, 163, 184)

    # 2. Document Title
    p_title = doc.add_paragraph()
    p_title.paragraph_format.space_before = Pt(0)
    p_title.paragraph_format.space_after = Pt(4)
    r_title = p_title.add_run("Adaptive Loop: Complete Interactive Demo Walkthrough")
    r_title.bold = True
    r_title.font.name = "Arial"
    r_title.font.size = Pt(24)
    r_title.font.color.rgb = RGBColor(15, 23, 42)

    # Subtitle
    p_sub = doc.add_paragraph()
    p_sub.paragraph_format.space_before = Pt(0)
    p_sub.paragraph_format.space_after = Pt(16)
    r_sub = p_sub.add_run("A Beginner-Friendly, Visual Guide to Clinical Decision Support for Longevity Medicine\nEmbedded Screenshots & Plain-English Explanations")
    r_sub.font.name = "Arial"
    r_sub.font.size = Pt(12)
    r_sub.font.color.rgb = RGBColor(71, 85, 105)

    add_callout(
        doc,
        "This document is written for everyone, especially those without any medical or computer science background. Every screen, button, metric, and AI thought process from the live application is explained in clear, everyday English with real embedded screenshots.",
        title="Executive Summary & Reading Guide",
        hex_bg="F0FDF4",
        hex_border="16A34A"
    )

    # -------------------------------------------------------------
    # SECTION 1
    # -------------------------------------------------------------
    h1 = doc.add_heading(level=1)
    r_h1 = h1.add_run("1. The Big Picture: Why Does Adaptive Loop Exist?")
    r_h1.font.name = "Arial"
    r_h1.font.size = Pt(16)
    r_h1.font.color.rgb = RGBColor(2, 132, 199)

    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.space_after = Pt(8)
    p.add_run(
        "Imagine you visit a preventative health clinic. You want to optimize your heart health and live a long, vibrant life. "
        "The doctor draws your blood, writes you a prescription for a supplement (such as Citrus Bergamot) and recommends that you do "
        "45 minutes of Zone 2 cardio exercise three times a week. Then the doctor smiles and says: "
    )
    r_quote = p.add_run('“See you in six months for your follow-up blood test!”\n\n')
    r_quote.bold = True
    p.add_run(
        "For the next 180 days, you live inside what medical researchers call the “Six-Month Black Box”. During those six months, you wear an Apple Watch, "
        "Whoop, or Oura Ring that measures your heart rate, sleep phases, and stress levels every single day. But your doctor never sees any of that data. "
        "When you finally return six months later, the doctor only looks at two data points: your Day 0 blood test and your Day 180 blood test."
    )

    add_callout(
        doc,
        "The Attribution Fallacy occurs when a doctor or algorithm looks at an improved blood test and blindly assumes the supplement caused it, "
        "ignoring that the patient may have quit a toxic job, changed their diet, or had a viral illness. Health is complex; isolated blood tests lie.",
        title="What is the 'Attribution Fallacy'?",
        hex_bg="FEF2F2",
        hex_border="EF4444"
    )

    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.space_after = Pt(8)
    p.add_run(
        "Adaptive Loop is an Intervention Intelligence Engine. It bridges this gap by acting like an expert clinical co-pilot. "
        "It continuously connects daily smartwatch biometrics (180 days of continuous data) with blood lab panels and tracks acute life events "
        "(such as respiratory illnesses or career resignations). It then applies a 4-stage neuro-symbolic reasoning pipeline—powered by Groq's "
        "openai/gpt-oss-120b model—to determine whether health protocols are genuinely working or being confounded by life."
    )

    # -------------------------------------------------------------
    # SECTION 2
    # -------------------------------------------------------------
    h1 = doc.add_heading(level=1)
    r_h1 = h1.add_run("2. The Four Core Principles (Made Simple)")
    r_h1.font.name = "Arial"
    r_h1.font.size = Pt(16)
    r_h1.font.color.rgb = RGBColor(2, 132, 199)

    table = doc.add_table(rows=5, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    widths = [Inches(1.8), Inches(2.2), Inches(2.5)]

    headers = ["Principle", "What It Means in Plain English", "Why It Matters to Patients"]
    for i, h in enumerate(headers):
        cell = table.cell(0, i)
        cell.width = widths[i]
        set_cell_background(cell, "0284C7")
        set_cell_margins(cell, top=120, bottom=120, left=140, right=140)
        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(h)
        r.bold = True
        r.font.name = "Arial"
        r.font.size = Pt(9.5)
        r.font.color.rgb = RGBColor(255, 255, 255)

    data = [
        ("Zero Algorithmic Dishonesty", "No fake mock data, no pre-calculated results stored in advance. Every analysis is computed live.", "You can trust that clinical recommendations represent actual real-time reasoning, never pre-written marketing scripts."),
        ("Neuro-Symbolic Architecture", "Combines AI intuition ('neuro') with exact mathematical calculations ('symbolic').", "Large language models can make arithmetic errors. Cold code computes the numbers; AI provides medical nuance."),
        ("Dual-Agent Adversarial Audit", "An Analyst AI writes the initial assessment, and a skeptical Critic AI audits it for errors.", "Two heads are better than one: the Critic prevents overconfidence and catches unearned credit."),
        ("Principled Abstention Guard", "If data is missing or compliance is under 40%, the system refuses to guess.", "In medicine, saying 'I do not have enough data to know' is infinitely safer than making a confident wild guess.")
    ]

    for row_idx, row_data in enumerate(data, start=1):
        bg = "F8FAFC" if row_idx % 2 == 1 else "FFFFFF"
        for col_idx, text in enumerate(row_data):
            cell = table.cell(row_idx, col_idx)
            cell.width = widths[col_idx]
            set_cell_background(cell, bg)
            set_cell_margins(cell, top=100, bottom=100, left=140, right=140)
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.15
            r = p.add_run(text)
            r.font.name = "Arial"
            r.font.size = Pt(9)
            if col_idx == 0:
                r.bold = True
                r.font.color.rgb = RGBColor(15, 23, 42)
            else:
                r.font.color.rgb = RGBColor(51, 65, 85)

    p_spacer = doc.add_paragraph()
    p_spacer.paragraph_format.space_after = Pt(12)

    # -------------------------------------------------------------
    # SECTION 3: Step 1 (Dashboard Landing)
    # -------------------------------------------------------------
    h1 = doc.add_heading(level=1)
    r_h1 = h1.add_run("3. Step 1: The Clinical Command Center (Dashboard Landing)")
    r_h1.font.name = "Arial"
    r_h1.font.size = Pt(16)
    r_h1.font.color.rgb = RGBColor(2, 132, 199)

    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.space_after = Pt(8)
    p.add_run(
        "Upon opening the application at http://localhost:3000, you are greeted by the Adaptive Loop Command Center. "
        "The interface is designed with a sleek, high-contrast dark theme optimized for clinical environments."
    )

    add_screenshot_figure(
        doc,
        "01_dashboard_landing.png",
        "Figure 1: The Adaptive Loop Command Center landing view, featuring cohort filters, benchmark patient cards, and staged un-computed telemetry."
    )

    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.space_after = Pt(8)
    p.add_run(
        "Key elements visible in Figure 1:\n"
        "• Top Navigation Bar: Displays the brand identity, active model indicator (openai/gpt-oss-120b), and a clinical decision support badge.\n"
        "• Clinical Cohort Filter: Enables sorting by preventative category (Cardiometabolic, Sleep Architecture, Metabolic Syndrome).\n"
        "• Patient Persona Cards: Highlights Marcus Vance (Clean Benchmark), Elena Rostova (Confounded Benchmark), and David Chen (Sparse Benchmark).\n"
        "• No Pre-Calculated Results: In adherence to zero algorithmic dishonesty, the dashboard starts with staged telemetry. Results are only computed when requested."
    )

    # -------------------------------------------------------------
    # SECTION 4: Step 2 (Marcus Vance Baseline)
    # -------------------------------------------------------------
    h1 = doc.add_heading(level=1)
    r_h1 = h1.add_run("4. Step 2: Exploring Patient Telemetry & Protocols (Marcus Vance)")
    r_h1.font.name = "Arial"
    r_h1.font.size = Pt(16)
    r_h1.font.color.rgb = RGBColor(2, 132, 199)

    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.space_after = Pt(8)
    p.add_run(
        "We click on Marcus Vance (MEM-1001), a 42-year-old male in the Cardiometabolic & Longevity cohort. "
        "Scrolling down reveals his detailed 180-day longitudinal biometric trajectory."
    )

    add_screenshot_figure(
        doc,
        "02_marcus_telemetry_protocols.png",
        "Figure 2: Marcus Vance's 180-day biometric trajectory chart, blood biomarker deltas, and active protocol adherence."
    )

    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.space_after = Pt(8)
    p.add_run(
        "Understanding Marcus's Health Metrics:\n"
        "• Apolipoprotein B (ApoB): Dropped from 118 mg/dL to 94 mg/dL (-20.3%). ApoB measures the concentration of atherogenic particles in the blood; this is a major cardiovascular risk reduction.\n"
        "• High-Sensitivity CRP: Plunged from 2.4 mg/L to 1.1 mg/L (-54.2%), indicating a steep drop in systemic blood vessel inflammation.\n"
        "• Wearable Trajectory: Over 26 weeks, his Heart Rate Variability (HRV - green line) climbed from 37.3 ms to 48.8 ms, while his Resting Heart Rate (red dotted line) dropped from 61.2 bpm to 55.6 bpm.\n"
        "• Protocol Adherence: Marcus logged an outstanding 88.3% adherence to his Zone 2 cardio and Citrus Bergamot regimen with zero confounding events."
    )

    # -------------------------------------------------------------
    # SECTION 5: Step 3 (Triggering Live Analysis)
    # -------------------------------------------------------------
    h1 = doc.add_heading(level=1)
    r_h1 = h1.add_run("5. Step 3: Triggering Live Multi-Stage AI Analysis")
    r_h1.font.name = "Arial"
    r_h1.font.size = Pt(16)
    r_h1.font.color.rgb = RGBColor(2, 132, 199)

    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.space_after = Pt(8)
    p.add_run(
        "We click the green “Run Adaptive Analysis” button. The interface immediately transitions into the live execution tracker, "
        "providing clinicians with transparent, stage-by-stage visibility into the neuro-symbolic reasoning pipeline."
    )

    add_screenshot_figure(
        doc,
        "03_marcus_analysis_executing.png",
        "Figure 3: The live neuro-symbolic execution tracker actively running inference via Groq's openai/gpt-oss-120b model."
    )

    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.space_after = Pt(8)
    p.add_run(
        "What happens during this live execution:\n"
        "1. Stage 1 (Data Aggregation): Python aggregates 180 individual daily records into 26 rolling weekly means and verifies data completeness.\n"
        "2. Stage 2 (AI Analyst): The structured profile is dispatched to Groq's openai/gpt-oss-120b model to formulate causal hypotheses connecting protocols to biomarker deltas.\n"
        "3. Stage 3 (Deterministic Citation Check): Cold code validates every observation ID cited by the AI against the SQLite observation registry to prevent hallucinations.\n"
        "4. Stage 4 (Adversarial Critic Audit): An adversarial AI Critic audits the Analyst's hypotheses against the verification report and lifestyle confounder timeline."
    )

    # -------------------------------------------------------------
    # SECTION 6: Step 4 (Clean Responder Synthesis)
    # -------------------------------------------------------------
    h1 = doc.add_heading(level=1)
    r_h1 = h1.add_run("6. Step 4: Clean Responder Synthesis (Marcus Vance)")
    r_h1.font.name = "Arial"
    r_h1.font.size = Pt(16)
    r_h1.font.color.rgb = RGBColor(2, 132, 199)

    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.space_after = Pt(8)
    p.add_run(
        "Within seconds, the completed clinical analysis dossier appears on screen, presenting the synthesized recommendation and dual-agent audit cards."
    )

    add_screenshot_figure(
        doc,
        "04_marcus_analysis_completed.png",
        "Figure 4: Marcus Vance's completed analysis report showing 'Continue Protocol As Prescribed', 100% grounded citations, and Critic approval."
    )

    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.space_after = Pt(8)
    p.add_run(
        "Key Elements of the Completed Report:\n"
        "• Canonical Action Badge: 'Continue Protocol As Prescribed' in green. Both Analyst and Critic agree the protocol is highly effective.\n"
        "• Primary Intervention Hypothesis: Formulated with moderate-to-high confidence, explicitly citing observation IDs OBS-BIO-APOB-D180, OBS-BIO-CRP-D180, and OBS-WEAR-HRV-W26.\n"
        "• Competing Hypotheses: Responsibly considers regression to the mean and background dietary improvements.\n"
        "• Deterministic Citations: 100% Grounded. Zero hallucinated data points.\n"
        "• Critic Verdict: APPROVED. Confirms the physiological response is genuine and unconfounded."
    )

    # -------------------------------------------------------------
    # SECTION 7: Step 5 (Confounder Detection - Elena Rostova)
    # -------------------------------------------------------------
    h1 = doc.add_heading(level=1)
    r_h1 = h1.add_run("7. Step 5: Confounder Detection & Attribution Fallacy (Elena Rostova)")
    r_h1.font.name = "Arial"
    r_h1.font.size = Pt(16)
    r_h1.font.color.rgb = RGBColor(2, 132, 199)

    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.space_after = Pt(8)
    p.add_run(
        "Now let us examine Elena Rostova (MEM-1002), a 36-year-old female in the Neuro-Endocrine & Sleep Architecture cohort. "
        "Her goal is deep sleep restoration and chronic stress reduction. Her morning cortisol dropped by -37.3% and fasting glucose dropped by -7.1%."
    )

    add_screenshot_figure(
        doc,
        "05_elena_confounder_timeline.png",
        "Figure 5: Elena Rostova's profile and confounder timeline displaying a high-stress career resignation (W08) and viral respiratory infection (W14)."
    )

    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.space_after = Pt(8)
    p.add_run(
        "Notice the Confounder Timeline in Figure 5:\n"
        "• CONF-CAREER-W08 (Severe Life Stress): In Week 8, Elena resigned from a grueling executive role. Her physical training load dropped 70% and sleep opportunity surged.\n"
        "• CONF-ILLNESS-W14 (Moderate Illness): In Week 14, an acute viral respiratory infection crashed her HRV and raised resting heart rate for six days.\n"
        "A naïve system would claim Magnesium L-Threonate caused the cortisol drop. Let us see how the Adversarial Critic handles this."
    )

    # -------------------------------------------------------------
    # SECTION 8: Step 6 (Elena's Adversarial Critic Audit)
    # -------------------------------------------------------------
    h1 = doc.add_heading(level=1)
    r_h1 = h1.add_run("8. Step 6: The AI Critic in Action (Elena's Adversarial Audit)")
    r_h1.font.name = "Arial"
    r_h1.font.size = Pt(16)
    r_h1.font.color.rgb = RGBColor(2, 132, 199)

    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.space_after = Pt(8)
    p.add_run(
        "We trigger live analysis for Elena Rostova. The results demonstrate the power of adversarial dual-agent auditing."
    )

    add_screenshot_figure(
        doc,
        "06_elena_analysis_critic_audit.png",
        "Figure 6: Elena Rostova's analysis results showing the Critic issuing 'VERDICT: AMENDED' due to the career transition confounder."
    )

    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.space_after = Pt(8)
    p.add_run(
        "The Critic's Intervention:\n"
        "• Initial Recommendation: CONTINUE_PROTOCOL.\n"
        "• Critic Verdict: VERDICT: AMENDED (Golden Amber Warning Badge).\n"
        "• Critic Evaluation: The Critic directly calls out that the career transition in Week 8 drastically reduced allostatic load and psychological cortisol drivers. Attributing the improvement solely to the supplement is an attribution error.\n"
        "• Clinical Action: The final action is amended to EXTEND_OBSERVATION so the clinician can observe Elena without confounder overlap."
    )

    # -------------------------------------------------------------
    # SECTION 9: Step 7 (David Chen Principled Abstention)
    # -------------------------------------------------------------
    h1 = doc.add_heading(level=1)
    r_h1 = h1.add_run("9. Step 7: Principled Clinical Abstention (David Chen)")
    r_h1.font.name = "Arial"
    r_h1.font.size = Pt(16)
    r_h1.font.color.rgb = RGBColor(2, 132, 199)

    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.space_after = Pt(8)
    p.add_run(
        "Next, we switch to David Chen (MEM-1003) in the Metabolic Syndrome & Glycemic Control cohort. "
        "David frequently forgot to wear his tracking sensor, creating large data gaps between Weeks 18 and 26. "
        "His logged protocol adherence is only 27.8% (indicated in the red adherence pill)."
    )

    add_screenshot_figure(
        doc,
        "07_david_abstention_guard.png",
        "Figure 7: David Chen's analysis triggering Principled Abstention ('Abstain / More Data Required') due to low adherence (27.8%) and sparse data."
    )

    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.space_after = Pt(8)
    p.add_run(
        "Safety Enforcement in Action:\n"
        "• Hard State Guard: If data completeness is under 50% or adherence is under 40%, the engine triggers an automatic safety halt.\n"
        "• Instant Return: The system returns in under 100 milliseconds without calling the LLM or generating speculative hypotheses.\n"
        "• Action Badge: 'Abstain / More Data Required' in purple.\n"
        "• Warning Message: Explains to the clinician that adherence was too low (27.8%) to draw reliable scientific conclusions."
    )

    # -------------------------------------------------------------
    # SECTION 10: Step 8 (Custom Patient & Cohort Modal)
    # -------------------------------------------------------------
    h1 = doc.add_heading(level=1)
    r_h1 = h1.add_run("10. Step 8: Adding a Custom Test Case Member & Cohort")
    r_h1.font.name = "Arial"
    r_h1.font.size = Pt(16)
    r_h1.font.color.rgb = RGBColor(2, 132, 199)

    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.space_after = Pt(8)
    p.add_run(
        "Clinicians and researchers can test their own cohorts and protocols. Clicking the '+ Add Test Case User' button opens an interactive modal."
    )

    add_screenshot_figure(
        doc,
        "08_add_member_modal.png",
        "Figure 8: The '+ Add Test Case User' modal filled out for Dr. Julian Croft in the 'Cellular Longevity & NAD+ Recovery' cohort."
    )

    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.space_after = Pt(8)
    p.add_run(
        "Configuring Dr. Julian Croft:\n"
        "• Full Name: Dr. Julian Croft | Age: 47 | Biological Sex: MALE.\n"
        "• Target Cohort: Cellular Longevity & NAD+ Recovery.\n"
        "• Health Goal: Optimize NAD+ synthesis & mitochondrial fatty acid oxidation.\n"
        "• Telemetry Profile: Clean Responder (generates 180 days of realistic daily biometrics with Gaussian noise).\n"
        "Clicking 'Add Member to Database' initiates real-time mathematical generation and commits the patient record to SQLite."
    )

    # -------------------------------------------------------------
    # SECTION 11: Step 9 (Custom Member Created)
    # -------------------------------------------------------------
    h1 = doc.add_heading(level=1)
    r_h1 = h1.add_run("11. Step 9: The Newly Created Member in Action")
    r_h1.font.name = "Arial"
    r_h1.font.size = Pt(16)
    r_h1.font.color.rgb = RGBColor(2, 132, 199)

    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.space_after = Pt(8)
    p.add_run(
        "The modal closes, and Dr. Julian Croft (MEM-2001) is immediately selected in the dashboard with fully synthesized longitudinal data."
    )

    add_screenshot_figure(
        doc,
        "09_custom_member_created.png",
        "Figure 9: Dr. Julian Croft loaded into the dashboard with 'Custom Test Case' badge, delete button, interactive tooltips, and un-computed analysis state."
    )

    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.space_after = Pt(8)
    p.add_run(
        "Key Highlights in Figure 9:\n"
        "• Custom Badge & Delete Icon: Clearly demarcates user-created test members and allows one-click deletion.\n"
        "• Interactive Hover Tooltip: Hovering over Week 20 shows exact biometrics (HRV: 42.8 ms, Resting HR: 62.8 bpm, Total Sleep: 7.13 hrs, Deep Sleep: 1.51 hrs).\n"
        "• Fresh Slate: Analysis has not been pre-computed. The green 'Run Adaptive Analysis' button awaits activation."
    )

    # -------------------------------------------------------------
    # SECTION 12: Step 10 (Custom Member Analysis Result)
    # -------------------------------------------------------------
    h1 = doc.add_heading(level=1)
    r_h1 = h1.add_run("12. Step 10: Running Live Analysis on the Custom Member")
    r_h1.font.name = "Arial"
    r_h1.font.size = Pt(16)
    r_h1.font.color.rgb = RGBColor(2, 132, 199)

    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.space_after = Pt(8)
    p.add_run(
        "We click 'Run Adaptive Analysis' for Dr. Julian Croft. The neuro-symbolic engine processes his profile live."
    )

    add_screenshot_figure(
        doc,
        "10_custom_member_analysis_result.png",
        "Figure 10: Live analysis results for custom member Dr. Julian Croft, confirming protocol continuation and grounded evidence citations."
    )

    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.space_after = Pt(8)
    p.add_run(
        "Live In-Flight Results for Dr. Croft:\n"
        "• Biomarker Changes: ApoB decreased by -21.1% (128 to 101 mg/dL) and hs-CRP decreased by -53.8% (2.6 to 1.2 mg/L).\n"
        "• Adherence: Confirmed at 86.1% over 180 days with zero acute confounders.\n"
        "• Action Recommendation: 'Continue Protocol As Prescribed' with Critic approval and 100% citation grounding."
    )

    # -------------------------------------------------------------
    # SECTION 13: Summary Table
    # -------------------------------------------------------------
    h1 = doc.add_heading(level=1)
    r_h1 = h1.add_run("13. Benchmark Scenario Comparison Matrix")
    r_h1.font.name = "Arial"
    r_h1.font.size = Pt(16)
    r_h1.font.color.rgb = RGBColor(2, 132, 199)

    summary_table = doc.add_table(rows=5, cols=5)
    summary_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    summary_table.autofit = False
    s_widths = [Inches(1.5), Inches(1.5), Inches(1.1), Inches(1.1), Inches(1.3)]

    s_headers = ["Persona", "Cohort", "Adherence", "Critic Verdict", "Final Clinical Action"]
    for i, h in enumerate(s_headers):
        cell = summary_table.cell(0, i)
        cell.width = s_widths[i]
        set_cell_background(cell, "0284C7")
        set_cell_margins(cell, top=100, bottom=100, left=120, right=120)
        p = cell.paragraphs[0]
        r = p.add_run(h)
        r.bold = True
        r.font.name = "Arial"
        r.font.size = Pt(9)
        r.font.color.rgb = RGBColor(255, 255, 255)

    s_data = [
        ("Marcus Vance (MEM-1001)", "Cardiometabolic & Longevity", "88.3%", "APPROVED", "CONTINUE_PROTOCOL"),
        ("Elena Rostova (MEM-1002)", "Neuro-Endocrine & Sleep", "62.2%", "AMENDED", "EXTEND_OBSERVATION"),
        ("David Chen (MEM-1003)", "Metabolic Syndrome", "27.8%", "N/A (Guard)", "ABSTAIN_MORE_DATA"),
        ("Dr. Julian Croft (MEM-2001)", "Cellular Longevity & NAD+", "86.1%", "APPROVED", "CONTINUE_PROTOCOL")
    ]

    for row_idx, row_data in enumerate(s_data, start=1):
        bg = "F8FAFC" if row_idx % 2 == 1 else "FFFFFF"
        for col_idx, text in enumerate(row_data):
            cell = summary_table.cell(row_idx, col_idx)
            cell.width = s_widths[col_idx]
            set_cell_background(cell, bg)
            set_cell_margins(cell, top=90, bottom=90, left=120, right=120)
            p = cell.paragraphs[0]
            r = p.add_run(text)
            r.font.name = "Arial"
            r.font.size = Pt(8.5)
            if col_idx == 0:
                r.bold = True
            if col_idx == 4:
                r.bold = True
                if "CONTINUE" in text:
                    r.font.color.rgb = RGBColor(22, 163, 74)
                elif "EXTEND" in text:
                    r.font.color.rgb = RGBColor(217, 119, 6)
                else:
                    r.font.color.rgb = RGBColor(147, 51, 234)

    p_spacer = doc.add_paragraph()
    p_spacer.paragraph_format.space_after = Pt(16)

    # -------------------------------------------------------------
    # SECTION 14: Reproduction Commands
    # -------------------------------------------------------------
    h1 = doc.add_heading(level=1)
    r_h1 = h1.add_run("14. Reproduction & Setup Instructions")
    r_h1.font.name = "Arial"
    r_h1.font.size = Pt(16)
    r_h1.font.color.rgb = RGBColor(2, 132, 199)

    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.space_after = Pt(8)
    p.add_run(
        "To run the live application or execute the automated Playwright demo and screenshot capture independently:\n\n"
        "1. Start the FastAPI backend server:\n"
        "   cd backend\n"
        "   python -m uvicorn app.main:app --host 0.0.0.0 --port 8000\n\n"
        "2. Start the Next.js frontend application:\n"
        "   cd frontend\n"
        "   npm run dev\n\n"
        "3. Run the automated interactive demo & screenshot capture:\n"
        "   python run_demo_and_capture.py"
    )

    # Save to both target locations
    os.makedirs(os.path.dirname(OUTPUT_DOCX_DOCS), exist_ok=True)
    os.makedirs(os.path.dirname(OUTPUT_DOCX_ARTIFACT), exist_ok=True)
    doc.save(OUTPUT_DOCX_DOCS)
    doc.save(OUTPUT_DOCX_ARTIFACT)
    print(f"Document successfully created and saved to:\n- {OUTPUT_DOCX_DOCS}\n- {OUTPUT_DOCX_ARTIFACT}")

if __name__ == "__main__":
    build_docx()
